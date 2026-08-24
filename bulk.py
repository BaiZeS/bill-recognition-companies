# 规则文件 Excel 解析/生成与批量导出文件组装。
import io
import json
import re
import zipfile
from datetime import datetime, timezone

from openpyxl import Workbook, load_workbook

from extract import validate_rules

RULE_HEADERS = ["字段编码", "中文名", "必填", "类型", "长度", "回填策略"]
_ROW_RE = re.compile(r"第 (\d+) 行")


def _cell_str(c):
    if c is None:
        return ""
    if isinstance(c, float) and c.is_integer():
        return str(int(c))
    return str(c).strip()


def parse_rules_xlsx(data: bytes):
    """解析规则 Excel。返回 (行字典列表, 对应工作表行号)；结构错误抛 ValueError。"""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True)
    except Exception as e:  # noqa: BLE001 - 任何坏文件统一报校验错误
        raise ValueError(f"Excel 解析失败：{e}") from None
    rows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise ValueError("Excel 为空")
    header = [_cell_str(c) for c in rows[0]][:6]
    if header != RULE_HEADERS:
        raise ValueError(f"表头应为：{' | '.join(RULE_HEADERS)}")
    rules, row_map = [], []
    for i, row in enumerate(rows[1:], start=2):
        cells = ([_cell_str(c) for c in row] + [""] * 6)[:6]
        if not any(cells):
            continue
        rules.append({
            "code": cells[0], "name": cells[1], "required": cells[2] or "可选",
            "type": cells[3], "length": cells[4] or None, "strategy": cells[5],
        })
        row_map.append(i)
    if not rules:
        raise ValueError("Excel 中没有规则行")
    return rules, row_map


def _remap_row_errors(msg: str, row_map: list[int]) -> str:
    def repl(m):
        idx = int(m.group(1)) - 1
        return f"第 {row_map[idx]} 行" if 0 <= idx < len(row_map) else m.group(0)
    return _ROW_RE.sub(repl, msg, count=1)


def load_rules_file(filename: str, data: bytes) -> list[dict]:
    """按扩展名解析 .xlsx/.json 规则文件，返回 validate_rules 规范化结果；失败抛 ValueError。"""
    name = filename.lower()
    if name.endswith(".xlsx"):
        rules, row_map = parse_rules_xlsx(data)
        try:
            return validate_rules(rules)
        except ValueError as e:
            raise ValueError(_remap_row_errors(str(e), row_map)) from None
    if name.endswith(".json"):
        try:
            payload = json.loads(data.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError) as e:
            raise ValueError(f"JSON 解析失败：{e}") from None
        if isinstance(payload, dict) and isinstance(payload.get("rules"), list):
            payload = payload["rules"]
        return validate_rules(payload)
    raise ValueError("仅支持 .xlsx 或 .json 规则文件")


def rules_to_xlsx(rules: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "规则"
    ws.append(RULE_HEADERS)
    for r in rules:
        ws.append([r["code"], r["name"], r["required"], r["type"],
                   r["length"] if r.get("length") is not None else "",
                   r.get("strategy", "")])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
