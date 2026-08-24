import io

from conftest import make_task
from openpyxl import Workbook, load_workbook

import bulk

XLSX_MT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def xlsx_bytes(rows):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_rules_seed_default(client):
    c, appmod = client
    tid, _ = make_task(appmod, "种子任务")
    resp = c.get(f"/api/tasks/{tid}/rules")
    assert resp.status_code == 200
    assert len(resp.json()["rules"]) == 21


def test_rules_seed_from_legacy_global(client, tmp_path):
    c, appmod = client
    (tmp_path / "rule_template.json").write_text(
        '{"rules": [{"code": "OLD1", "name": "旧字段", "required": "必填",'
        ' "type": "字符", "length": null, "strategy": ""}]}',
        encoding="utf-8")
    tid, _ = make_task(appmod, "迁移任务")
    rules = c.get(f"/api/tasks/{tid}/rules").json()["rules"]
    assert [r["code"] for r in rules] == ["OLD1"]


def test_put_rules_persists(client):
    c, appmod = client
    tid, _ = make_task(appmod, "保存规则")
    payload = {"rules": [{"code": "A1", "name": "甲", "required": "必填",
                          "type": "字符", "length": None, "strategy": ""}]}
    assert c.put(f"/api/tasks/{tid}/rules", json=payload).status_code == 200
    assert (appmod.TASKS_DIR / tid / "rules.json").exists()
    appmod.RULES.clear()  # 强制走文件读取
    assert c.get(f"/api/tasks/{tid}/rules").json()["rules"][0]["code"] == "A1"


def test_put_rules_duplicate_code(client):
    c, appmod = client
    tid, _ = make_task(appmod, "重复编码")
    rule = {"code": "A1", "name": "甲", "required": "必填", "type": "字符",
            "length": None, "strategy": ""}
    resp = c.put(f"/api/tasks/{tid}/rules", json={"rules": [rule, dict(rule)]})
    assert resp.status_code == 422
    assert "第 2 行" in resp.json()["detail"]


def test_upload_rules_xlsx(client):
    c, appmod = client
    tid, _ = make_task(appmod, "上传规则")
    data = xlsx_bytes([bulk.RULE_HEADERS, ["U1", "上传字段", "必填", "字符", "", ""]])
    resp = c.post(f"/api/tasks/{tid}/rules/upload",
                  files={"file": ("r.xlsx", data, XLSX_MT)})
    assert resp.status_code == 200
    assert c.get(f"/api/tasks/{tid}/rules").json()["rules"][0]["code"] == "U1"


def test_upload_rules_bad_header(client):
    c, appmod = client
    tid, _ = make_task(appmod, "坏表头")
    data = xlsx_bytes([["编码"], ["x"]])
    resp = c.post(f"/api/tasks/{tid}/rules/upload",
                  files={"file": ("r.xlsx", data, XLSX_MT)})
    assert resp.status_code == 422


def test_download_rules_xlsx(client):
    c, appmod = client
    tid, _ = make_task(appmod, "下载规则")
    resp = c.get(f"/api/tasks/{tid}/rules.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws["A1"].value == "字段编码"
    assert ws.max_row == 22  # 表头 + 21 默认字段


def test_rules_default_endpoints(client):
    c, _ = client
    assert len(c.get("/api/rules-default").json()["rules"]) == 21
    assert c.get("/api/rules-default.xlsx").status_code == 200


def test_global_template_endpoints_removed(client):
    c, _ = client
    assert c.get("/api/rule-template").status_code == 404


def test_extract_uses_task_rules(client, monkeypatch):
    c, appmod = client
    monkeypatch.setenv("BAILIAN_API_KEY", "test-key")
    calls = []

    def fake_llm(system, user, api_key, base_url, model):
        calls.append((system, user))
        return '{"fields": {}}'

    monkeypatch.setattr(appmod, "_call_llm", fake_llm)
    tid, _ = make_task(appmod, "按任务规则抽取")
    c.put(f"/api/tasks/{tid}/rules", json={"rules": [
        {"code": "MYFIELD", "name": "我的字段", "required": "必填",
         "type": "字符", "length": None, "strategy": ""}]})
    resp = c.post(f"/api/tasks/{tid}/extract", json={})
    assert resp.status_code == 200
    assert any("MYFIELD" in s + u for s, u in calls)
    assert (appmod.TASKS_DIR / tid / "extracted.json").exists()


def test_extract_payload_rules_override(client, monkeypatch):
    c, appmod = client
    monkeypatch.setenv("BAILIAN_API_KEY", "test-key")
    calls = []

    def fake_llm(system, user, api_key, base_url, model):
        calls.append((system, user))
        return '{"fields": {}}'

    monkeypatch.setattr(appmod, "_call_llm", fake_llm)
    tid, _ = make_task(appmod, "覆盖规则")
    resp = c.post(f"/api/tasks/{tid}/extract", json={"rules": [
        {"code": "OVR1", "name": "覆盖字段", "required": "必填",
         "type": "字符", "length": None, "strategy": ""}]})
    assert resp.status_code == 200
    assert any("OVR1" in s + u for s, u in calls)


def test_extract_rules_from_file_after_restart(client, monkeypatch):
    c, appmod = client
    monkeypatch.setenv("BAILIAN_API_KEY", "test-key")
    calls = []

    def fake_llm(system, user, api_key, base_url, model):
        calls.append((system, user))
        return '{"fields": {}}'

    monkeypatch.setattr(appmod, "_call_llm", fake_llm)
    tid, _ = make_task(appmod, "文件规则")
    c.put(f"/api/tasks/{tid}/rules", json={"rules": [
        {"code": "FILE1", "name": "文件字段", "required": "必填",
         "type": "字符", "length": None, "strategy": ""}]})
    appmod.RULES.clear()  # 模拟进程重启后缓存丢失
    resp = c.post(f"/api/tasks/{tid}/extract", json={})
    assert resp.status_code == 200
    assert any("FILE1" in s + u for s, u in calls)


def test_corrupted_rules_file_falls_back_to_seed(client):
    c, appmod = client
    tid, _ = make_task(appmod, "损坏规则")
    (appmod.TASKS_DIR / tid / "rules.json").write_text(
        '{"rules": [{"name": "无编码"}]}', encoding="utf-8")
    resp = c.get(f"/api/tasks/{tid}/rules")
    assert resp.status_code == 200
    assert len(resp.json()["rules"]) == 21  # 回退默认种子
