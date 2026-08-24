import io
import json
import zipfile

from conftest import make_task
from openpyxl import load_workbook

EXTRACTED_SAMPLE = {
    "model": "m1",
    "groups": [{
        "page": 0, "index": 0,
        "fields": [
            {"code": "tran_amt", "name": "金额", "value": "100.00",
             "required": "必填", "source_ids": []},
            {"code": "bankseqno", "name": "银行交易流水号", "value": "N1",
             "required": "必填", "source_ids": []},
        ],
    }],
}


def zip_names(body):
    return set(zipfile.ZipFile(io.BytesIO(body)).namelist())


def zip_read(body, name):
    return zipfile.ZipFile(io.BytesIO(body)).read(name)


def test_export_result_zip(client):
    c, appmod = client
    t1, _ = make_task(appmod, "导出甲")
    t2, _ = make_task(appmod, "导出乙")
    resp = c.post("/api/export", json={"task_ids": [t1, t2], "kind": "result"})
    assert resp.status_code == 200
    assert zip_names(resp.content) == {
        "导出甲_structured.json", "导出乙_structured.json", "识别结果汇总.xlsx"}
    ws = load_workbook(io.BytesIO(zip_read(resp.content, "识别结果汇总.xlsx"))).active
    assert [r[0] for r in ws.iter_rows(min_row=2, values_only=True)] == ["导出甲.pdf", "导出乙.pdf"]


def test_export_result_prefers_corrected(client):
    c, appmod = client
    tid, doc = make_task(appmod, "修正任务", has_corrected=True)
    corrected = [dict(doc[0])]
    corrected[0]["kv"] = [dict(doc[0]["kv"][0], value="999")]
    appmod.RESULTS[tid] = corrected
    resp = c.post("/api/export", json={"task_ids": [tid], "kind": "result"})
    body = resp.content
    assert "修正任务_corrected.json" in zip_names(body)
    saved = json.loads(zip_read(body, "修正任务_corrected.json"))
    assert saved[0]["kv"][0]["value"] == "999"


def test_export_extracted_zip_and_skip(client):
    c, appmod = client
    t1, _ = make_task(appmod, "有抽取", extracted=EXTRACTED_SAMPLE)
    t2, _ = make_task(appmod, "无抽取")
    resp = c.post("/api/export", json={"task_ids": [t1, t2], "kind": "extracted"})
    assert resp.status_code == 200
    assert resp.headers["X-Skipped"] == "1"
    flat = json.loads(zip_read(resp.content, "有抽取_extracted.json"))
    assert flat["_meta"]["filename"] == "有抽取.pdf"
    assert flat["results"] == [{"tran_amt": "100.00", "bankseqno": "N1"}]
    ws = load_workbook(io.BytesIO(zip_read(resp.content, "抽取结果汇总.xlsx"))).active
    header = [c2.value for c2 in ws[1]]
    assert header == ["文件名", "组序号", "页码", "金额", "银行交易流水号"]
    row = [c2.value for c2 in ws[2]]
    assert row == ["有抽取.pdf", 1, 1, "100.00", "N1"]


def test_export_extracted_all_skipped(client):
    c, appmod = client
    tid, _ = make_task(appmod, "都没抽取")
    resp = c.post("/api/export", json={"task_ids": [tid], "kind": "extracted"})
    assert resp.status_code == 422


def test_export_no_done_tasks(client):
    c, appmod = client
    resp = c.post("/api/export", json={"task_ids": ["nonexist"], "kind": "result"})
    assert resp.status_code == 422


def test_tasks_list_has_extracted(client):
    c, appmod = client
    t1, _ = make_task(appmod, "未抽取")
    t2, _ = make_task(appmod, "已抽取", extracted=EXTRACTED_SAMPLE)
    tasks = {t["id"]: t for t in c.get("/api/tasks").json()["tasks"]}
    assert tasks[t1]["has_extracted"] is False
    assert tasks[t2]["has_extracted"] is True
